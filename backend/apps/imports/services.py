import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from apps.customers.models import Customer

EXPECTED_HEADERS = [
    'Customer ID',
    'Customer Name',
    'Phone Number',
    'No. of Churn Days',
    'Package'
]

def generate_excel_template() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customer Upload Template"

    # Header styling
    header_fill = PatternFill(start_color="F44336", end_color="F44336", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    for col_idx, header in enumerate(EXPECTED_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Sample rows
    samples = [
        ["DH10021", "Ram Sharma", "9801234567", 95, "200 Mbps"],
        ["DH10022", "Sita Thapa", "9812345678", 120, "100 Mbps"],
        ["DH10023", "Hari Gurung", "9841234567", 180, "300 Mbps"],
    ]

    for row_idx, row_data in enumerate(samples, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx == 4:
                cell.alignment = Alignment(horizontal="right")

    # Column widths
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def parse_and_validate_file(file_obj, filename: str):
    """
    Parses an uploaded Excel (.xlsx) or CSV file.
    Validates headers and each row against rules:
    - Customer ID: required string
    - Customer Name: required
    - Phone Number: required string
    - No. of Churn Days: non-negative integer
    - Package: required string
    """
    if filename.lower().endswith('.csv'):
        import csv
        content = file_obj.read()
        if isinstance(content, bytes):
            text = content.decode('utf-8-sig', errors='ignore')
        else:
            text = str(content)
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    else:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

    if not rows:
        raise ValueError("Uploaded file is completely empty.")

    header_row = [str(cell).strip() if cell is not None else "" for cell in rows[0]]

    # Validate header matches expected columns (fuzzy or exact)
    header_map = {}
    for idx, name in enumerate(header_row):
        normalized = name.lower().replace(" ", "").replace("_", "").replace(".", "")
        if "customerid" in normalized or "custid" in normalized:
            header_map['customer_id'] = idx
        elif "customername" in normalized or "name" in normalized:
            header_map['customer_name'] = idx
        elif "phone" in normalized or "contact" in normalized or "mobile" in normalized:
            header_map['phone_number'] = idx
        elif "churn" in normalized or "days" in normalized:
            header_map['churn_days'] = idx
        elif "package" in normalized or "plan" in normalized:
            header_map['package'] = idx

    missing = []
    for req in ['customer_id', 'customer_name', 'phone_number', 'churn_days', 'package']:
        if req not in header_map:
            missing.append(req)

    if missing:
        raise ValueError(f"Uploaded template is missing required columns: {', '.join(missing)}. Please use the standard 5-column template.")

    existing_customer_ids = set(
        Customer.objects.values_list('customer_id', flat=True)
    )

    seen_file_ids = set()
    valid_rows = []
    errors = []
    duplicate_rows = []

    for row_num, row_cells in enumerate(rows[1:], start=2):
        if not any(row_cells):
            continue  # skip completely blank lines

        # Extract values
        cid_raw = row_cells[header_map['customer_id']] if header_map['customer_id'] < len(row_cells) else None
        name_raw = row_cells[header_map['customer_name']] if header_map['customer_name'] < len(row_cells) else None
        phone_raw = row_cells[header_map['phone_number']] if header_map['phone_number'] < len(row_cells) else None
        churn_raw = row_cells[header_map['churn_days']] if header_map['churn_days'] < len(row_cells) else None
        pkg_raw = row_cells[header_map['package']] if header_map['package'] < len(row_cells) else None

        customer_id = str(cid_raw).strip() if cid_raw is not None else ""
        customer_name = str(name_raw).strip() if name_raw is not None else ""
        phone_number = str(phone_raw).strip() if phone_raw is not None else ""
        package = str(pkg_raw).strip() if pkg_raw is not None else ""

        row_errors = []
        if not customer_id:
            row_errors.append("Customer ID is required.")
        if not customer_name:
            row_errors.append("Customer Name is required.")
        if not phone_number:
            row_errors.append("Phone Number is required.")
        if not package:
            row_errors.append("Package is required.")

        # Parse churn days
        churn_days = 0
        try:
            churn_days = int(float(churn_raw)) if churn_raw is not None else 0
            if churn_days < 0:
                row_errors.append("No. of Churn Days cannot be negative.")
        except (ValueError, TypeError):
            row_errors.append(f"Invalid churn days value: '{churn_raw}'. Must be a whole number.")

        is_file_duplicate = False
        is_db_duplicate = False

        if customer_id:
            if customer_id in seen_file_ids:
                is_file_duplicate = True
                row_errors.append("Duplicate Customer ID within the same file.")
            else:
                seen_file_ids.add(customer_id)

            if customer_id in existing_customer_ids:
                is_db_duplicate = True

        row_data = {
            'row_number': row_num,
            'customer_id': customer_id,
            'customer_name': customer_name,
            'phone_number': phone_number,
            'churn_days': churn_days,
            'package': package,
            'is_duplicate': is_file_duplicate or is_db_duplicate,
            'is_db_duplicate': is_db_duplicate,
            'errors': row_errors,
            'is_valid': len(row_errors) == 0
        }

        if row_errors:
            errors.append(row_data)
        elif is_db_duplicate or is_file_duplicate:
            duplicate_rows.append(row_data)
            valid_rows.append(row_data)
        else:
            valid_rows.append(row_data)

    total_rows = len(valid_rows) + len(errors)

    return {
        'total_rows': total_rows,
        'valid_count': len(valid_rows),
        'invalid_count': len(errors),
        'duplicate_count': len(duplicate_rows),
        'preview_rows': (valid_rows + errors)[:100],  # preview first 100 rows
        'all_valid_rows': valid_rows,
        'errors': errors,
        'duplicate_rows': duplicate_rows
    }
