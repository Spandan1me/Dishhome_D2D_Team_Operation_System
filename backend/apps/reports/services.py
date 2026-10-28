import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.db.models import Count, Q, Avg
from django.utils import timezone
from datetime import timedelta

from apps.customers.models import Customer, RecordSource, ReviewStatus
from apps.allocations.models import CustomerAssignment
from apps.visits.models import VisitLog, VisitStatus, FOLLOW_UP_REMARKS, SWITCHING_REASONS, SWITCHING_TO_ISPS
from apps.followups.models import FollowUp, FollowUpStatus
from apps.accounts.models import User, RoleChoices

def get_supervisor_dashboard_stats(user):
    today = timezone.localdate()

    # Supervisors need shared visibility across the operation workspace, not only
    # agents under their own direct supervision.
    team_agents = User.objects.filter(role=RoleChoices.AGENT) if not user.is_admin_user else User.objects.filter(role=RoleChoices.AGENT)

    # Customers scope
    if user.is_admin_user or user.is_supervisor_user:
        customers_qs = Customer.objects.all()
        visits_qs = VisitLog.objects.all()
        followups_qs = FollowUp.objects.all()
    else:
        customers_qs = Customer.objects.filter(
            Q(assignments__assigned_agent__in=team_agents) |
            Q(created_by=user) |
            Q(created_by__in=team_agents) |
            Q(batch_records__batch__uploaded_by=user) |
            Q(assignments__isnull=True)
        ).distinct()
        visits_qs = VisitLog.objects.filter(agent__in=team_agents)
        followups_qs = FollowUp.objects.filter(assigned_agent__in=team_agents)

    total_customers = customers_qs.count()
    total_assigned = customers_qs.filter(assignments__status='ACTIVE').distinct().count()
    unassigned_count = total_customers - total_assigned

    # Unique customers visited
    unique_customers_visited = visits_qs.filter(
        visit_status=VisitStatus.COMPLETED
    ).values('customer').distinct().count()

    completed_visits_count = visits_qs.filter(visit_status=VisitStatus.COMPLETED).count()
    pending_visits_count = visits_qs.filter(visit_status=VisitStatus.PENDING).count()

    followups_due_today = followups_qs.filter(
        follow_up_date=today,
        status=FollowUpStatus.PENDING
    ).count()

    overdue_followups = followups_qs.filter(
        Q(status=FollowUpStatus.OVERDUE) |
        Q(status=FollowUpStatus.PENDING, follow_up_date__lt=today)
    ).count()

    agent_added_pending = Customer.objects.filter(
        record_source=RecordSource.AGENT_MANUAL,
        review_status=ReviewStatus.PENDING_REVIEW
    ).count()

    gps_captured_visits = visits_qs.filter(
        gps_latitude__isnull=False,
        gps_longitude__isnull=False
    ).count()

    missing_gps_visits = visits_qs.filter(
        gps_latitude__isnull=True
    ).count()

    # Charts data
    # 1. Daily completion trend (last 14 days)
    daily_trend = []
    for i in range(13, -1, -1):
        day = today - timedelta(days=i)
        cnt = visits_qs.filter(
            completed_at__date=day,
            visit_status=VisitStatus.COMPLETED
        ).count()
        daily_trend.append({
            'date': day.strftime('%b %d'),
            'visits': cnt
        })

    # 2. Agent performance
    agent_performance = []
    for agent in team_agents.filter(is_active=True):
        assigned = CustomerAssignment.objects.filter(assigned_agent=agent, status='ACTIVE').count()
        completed = VisitLog.objects.filter(agent=agent, visit_status=VisitStatus.COMPLETED).count()
        agent_performance.append({
            'name': agent.full_name or agent.username,
            'assigned': assigned,
            'completed': completed
        })

    # 3. Follow-up remarks distribution (top 8)
    remark_counts = (
        visits_qs.values('follow_up_remark')
        .annotate(count=Count('id'))
        .order_by('-count')[:8]
    )
    remark_dist = [
        {'name': item['follow_up_remark'], 'value': item['count']}
        for item in remark_counts if item['follow_up_remark']
    ]

    # 4. Switching reasons distribution
    reason_counts = (
        visits_qs.exclude(switching_reason='')
        .values('switching_reason')
        .annotate(count=Count('id'))
        .order_by('-count')[:8]
    )
    reason_dist = [
        {'name': item['switching_reason'], 'value': item['count']}
        for item in reason_counts
    ]

    # 5. Competitor ISP distribution
    isp_counts = (
        visits_qs.exclude(switching_to='')
        .values('switching_to')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    isp_dist = [
        {'name': item['switching_to'], 'value': item['count']}
        for item in isp_counts
    ]

    return {
        'kpis': {
            'total_customers': total_customers,
            'total_assigned': total_assigned,
            'unassigned_count': unassigned_count,
            'unique_customers_visited': unique_customers_visited,
            'completed_visits_count': completed_visits_count,
            'pending_visits_count': pending_visits_count,
            'followups_due_today': followups_due_today,
            'overdue_followups': overdue_followups,
            'agent_added_pending': agent_added_pending,
            'gps_captured_visits': gps_captured_visits,
            'missing_gps_visits': missing_gps_visits,
        },
        'charts': {
            'daily_trend': daily_trend,
            'agent_performance': agent_performance,
            'remark_distribution': remark_dist,
            'reason_distribution': reason_dist,
            'isp_distribution': isp_dist
        }
    }

def get_agent_dashboard_stats(user):
    today = timezone.localdate()

    assigned_qs = CustomerAssignment.objects.filter(assigned_agent=user, status='ACTIVE')
    total_assigned = assigned_qs.count()

    visits_qs = VisitLog.objects.filter(agent=user)
    completed_visits = visits_qs.filter(visit_status=VisitStatus.COMPLETED).count()
    pending_visits = total_assigned - visits_qs.filter(
        visit_status=VisitStatus.COMPLETED
    ).values('customer').distinct().count()
    if pending_visits < 0:
        pending_visits = 0

    followups_qs = FollowUp.objects.filter(assigned_agent=user)
    followups_today = followups_qs.filter(
        follow_up_date=today,
        status=FollowUpStatus.PENDING
    ).count()
    overdue_followups = followups_qs.filter(
        Q(status=FollowUpStatus.OVERDUE) |
        Q(status=FollowUpStatus.PENDING, follow_up_date__lt=today)
    ).count()

    agent_added = Customer.objects.filter(
        created_by=user,
        record_source=RecordSource.AGENT_MANUAL
    ).count()

    return {
        'total_assigned': total_assigned,
        'completed_visits': completed_visits,
        'pending_visits': pending_visits,
        'followups_today': followups_today,
        'overdue_followups': overdue_followups,
        'agent_added_count': agent_added
    }

def create_excel_report(sheet_title: str, headers: list, rows_data: list) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    header_fill = PatternFill(start_color="F44336", end_color="F44336", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, row in enumerate(rows_data, 2):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            # Treat customer ID or phone number as string
            if col_idx in [1, 3]:
                cell.number_format = '@'

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
